"""Shared utilities, configuration, LLM invocation, memory, and document parsing for NextRole agents."""

import os
import re
import io
import json
import time
from typing import Dict, Any, List, Optional
from pathlib import Path
from dotenv import load_dotenv
from google import genai
from google.genai import types

load_dotenv()

TAVILY_API_KEY      = os.getenv("TAVILY_API_KEY", "")
LLAMA_CLOUD_API_KEY = os.getenv("LLAMA_CLOUD_API_KEY", "")
GOOGLE_API_KEY      = os.getenv("GOOGLE_API_KEY", "")
ANTHROPIC_API_KEY   = os.getenv("ANTHROPIC_API_KEY", "")
RAPIDAPI_KEY        = os.getenv("RAPIDAPI_KEY", "")

PRIMARY_MODEL   = "gemini-3.5-flash-lite"
FALLBACK_MODELS = ["gemini-3.8-flash", "gemini-flash-latest", "gemini-flash-lite-latest", "gemini-2.5-flash"]

MEMORY_PATH = Path(__file__).parent.parent / "career_memory.json"
MAX_REPLAN_LOOPS = 2


def get_genai_client() -> genai.Client:
    api_key = os.getenv("GOOGLE_API_KEY", "")
    if not api_key:
        raise ValueError("GOOGLE_API_KEY is not set in .env")
    return genai.Client(api_key=api_key)


def generate_llm(
    prompt: str,
    system_instruction: Optional[str] = None,
    model_override: Optional[str] = None
) -> str:
    """Invokes Gemini with automatic model failover and retries."""
    client = get_genai_client()
    models = ([model_override] + FALLBACK_MODELS) if model_override else ([PRIMARY_MODEL] + FALLBACK_MODELS)
    last_err = None

    for m in models:
        try:
            config = types.GenerateContentConfig(
                system_instruction=system_instruction
            ) if system_instruction else None
            resp = client.models.generate_content(model=m, contents=prompt, config=config)
            if resp and resp.text:
                return resp.text.strip()
        except Exception as e:
            last_err = e
            err_str = str(e).lower()
            print(f"[Model attempt failed] {m}: {e}")
            if any(k in err_str for k in ("quota", "resource_exhausted", "404", "not_found")):
                continue
            time.sleep(0.5)

    raise RuntimeError(f"All LLM models failed. Last error: {last_err}")


def extract_urls(text: str) -> List[str]:
    pattern = r"https?://(?:www\.)?[-a-zA-Z0-9@:%._\+~#=]{1,256}\.[a-zA-Z0-9()]{1,6}\b(?:[-a-zA-Z0-9()@:%_\+.~#?&//=]*)"
    return re.findall(pattern, text)


# ── CAREER MEMORY (persistent JSON) ──────────────────────────────────────────

def load_memory() -> Dict[str, Any]:
    if MEMORY_PATH.exists():
        try:
            return json.loads(MEMORY_PATH.read_text(encoding="utf-8"))
        except Exception:
            pass
    return {"profiles": {}, "sessions": []}


def save_memory(memory: Dict[str, Any]) -> None:
    MEMORY_PATH.write_text(json.dumps(memory, indent=2, ensure_ascii=False), encoding="utf-8")


def upsert_memory(candidate_id: str, artifacts: Dict[str, str], meta: Dict[str, Any] = None) -> None:
    memory = load_memory()
    profile = memory["profiles"].get(candidate_id, {"history": []})
    entry = {
        "timestamp": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime()),
        "artifacts": {k: v[:500] for k, v in artifacts.items()},
        "meta": meta or {}
    }
    profile["history"].append(entry)
    profile["latest_artifacts"] = {k: v[:3000] for k, v in artifacts.items()}
    memory["profiles"][candidate_id] = profile
    memory["sessions"].append({"id": candidate_id, "ts": entry["timestamp"]})
    save_memory(memory)


def get_memory_for_candidate(candidate_id: str) -> Optional[Dict[str, Any]]:
    memory = load_memory()
    return memory["profiles"].get(candidate_id)


# ── SHARED CAREER EVIDENCE STORE (session state) ─────────────────────────────

def get_evidence_store(session: Dict[str, Any]) -> Dict[str, Any]:
    """Returns the shared in-session Career Evidence Store linking all 4 feature teams."""
    if "evidence_store" not in session:
        session["evidence_store"] = {
            "skills": {},        # skill -> {strength: strong|partial|gap|unknown, evidence: str}
            "overlaps": [],      # proven match points
            "gaps": [],          # gap items with severity
            "learning_plan": {}, # day plans per gap topic
            "sim_history": [],   # interview sim Q&A log
            "job_results": [],   # normalized job cards from Job Finder
        }
    return session["evidence_store"]


# ── DOCUMENT PROCESSING ──────────────────────────────────────────────────────

def extract_docx_bytes(file_bytes: bytes) -> str:
    try:
        import docx
        doc = docx.Document(io.BytesIO(file_bytes))
        sections = []
        for p in doc.paragraphs:
            txt = p.text.strip()
            if txt:
                sections.append(txt)
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                dedup = []
                for c in cells:
                    if not dedup or dedup[-1] != c:
                        dedup.append(c)
                if dedup:
                    sections.append(" | ".join(dedup))
        combined = "\n\n".join(sections)
        if combined.strip():
            return combined.strip()
    except Exception as e:
        print(f"[Warning] python-docx error: {e}")

    try:
        import zipfile, xml.etree.ElementTree as ET
        with zipfile.ZipFile(io.BytesIO(file_bytes)) as z:
            xml_content = z.read("word/document.xml")
            tree = ET.fromstring(xml_content)
            paragraphs = []
            NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
            for p in tree.iter(NS + "p"):
                texts = [node.text for node in p.iter(NS + "t") if node.text]
                if texts:
                    val = "".join(texts).strip()
                    if val:
                        paragraphs.append(val)
            if paragraphs:
                return "\n\n".join(paragraphs)
    except Exception as e:
        print(f"[Warning] docx xml fallback error: {e}")

    return ""


def parse_resume_bytes(filename: str, file_bytes: bytes) -> str:
    lower_name = filename.lower()
    text = ""

    if lower_name.endswith(".docx") or lower_name.endswith(".doc"):
        text = extract_docx_bytes(file_bytes)
        if not text.strip() and LLAMA_CLOUD_API_KEY:
            try:
                from llama_parse import LlamaParse
                parser = LlamaParse(api_key=LLAMA_CLOUD_API_KEY, result_type="markdown", verbose=False)
                docs = parser.load_data(file_bytes, extra_info={"file_name": filename})
                if docs:
                    text = "\n\n".join([d.text for d in docs])
            except Exception as e:
                print(f"[Warning] LlamaParse DOCX error: {e}")

    elif lower_name.endswith(".pdf"):
        # 1. Fast local PDF extraction first (instant and reliable)
        try:
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            pages = [f"## Page {i+1}\n\n{p.extract_text()}" for i, p in enumerate(reader.pages) if p.extract_text()]
            text = "\n\n".join(pages).strip()
        except Exception as e:
            print(f"[Warning] pypdf extraction error: {e}")

        # 2. If local extraction was empty (e.g. scanned image PDF) and LlamaParse key is available
        if not text.strip() and LLAMA_CLOUD_API_KEY:
            try:
                from llama_parse import LlamaParse
                parser = LlamaParse(api_key=LLAMA_CLOUD_API_KEY, result_type="markdown", verbose=False)
                docs = parser.load_data(file_bytes, extra_info={"file_name": filename})
                if docs:
                    text = "\n\n".join([d.text for d in docs]).strip()
            except Exception as e:
                print(f"[Warning] LlamaParse PDF fallback error: {e}")

    if not text.strip():
        try:
            text = file_bytes.decode("utf-8", errors="ignore").strip()
        except Exception:
            text = ""
    return text.strip()


def extract_job_from_url(url: str) -> Dict[str, Any]:
    if not TAVILY_API_KEY:
        raise ValueError("TAVILY_API_KEY is missing in .env")
    from tavily import TavilyClient
    client = TavilyClient(api_key=TAVILY_API_KEY)
    raw_text = ""

    try:
        if hasattr(client, "extract"):
            res = client.extract(urls=[url])
            if res and "results" in res and res["results"]:
                raw_text = res["results"][0].get("raw_content", "")
    except Exception as e:
        print(f"[Notice] Tavily extract: {e}")

    if not raw_text:
        try:
            sr = client.search(query=f"job opening details {url}", max_results=3)
            raw_text = " ".join([r.get("content", "") for r in sr.get("results", [])])
        except Exception as e:
            print(f"[Notice] Tavily search fallback: {e}")

    if not raw_text.strip():
        raw_text = f"Job listing from {url}. Please review the web source."

    prompt = (
        "You are an expert technical recruiter. Structure this raw job description into clean Markdown.\n"
        f"Raw content:\n\"\"\"\n{raw_text[:9000]}\n\"\"\"\n"
        "Format:\n# <Company> - <Role>\n- **Location:** ...\n- **Employment Type:** ...\n- **Level:** ...\n"
        "## Role Summary\n## Key Responsibilities\n## Required Qualifications\n## Preferred Qualifications\n## Tech Stack"
    )
    try:
        md = generate_llm(prompt, system_instruction="Extract accurate, actionable job posting information.")
    except Exception as e:
        print(f"[Warning] JD LLM error: {e}")
        md = f"# Target Job Posting\n\n**Source URL:** {url}\n\n## Job Description\n\n{raw_text[:7000]}"
    return {"url": url, "markdown": md}
